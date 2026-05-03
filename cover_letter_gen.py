import streamlit as st
import os
import json
from groq import Groq
from dotenv import load_dotenv

from docx import Document
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet

from io import BytesIO

load_dotenv()
client = Groq(api_key=os.getenv("GROQ_API_KEY"))

with open("config.json", "r") as f:
    config = json.load(f)

st.set_page_config(page_title="Cover Letter Generator", layout="wide")
st.title("Cover Letter Generator")

def create_docx_bytes(text: str) -> BytesIO:
    doc = Document()
    doc.add_heading("Cover Letter", level=1)
    doc.add_paragraph(text)

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


def create_pdf_bytes(text: str) -> BytesIO:
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer)

    styles = getSampleStyleSheet()
    content = []

    for line in text.split("\n"):
        content.append(Paragraph(line, styles["Normal"]))
        content.append(Spacer(1, 8))

    doc.build(content)
    buffer.seek(0)
    return buffer


def generate():
    if not st.session_state.job_description or not st.session_state.resume:
        st.warning("Fill both inputs.")
        return

    with st.spinner("Generating..."):
        response = client.chat.completions.create(
            model=st.session_state.model,
            messages=[
                {"role": "system", "content": st.session_state.system_prompt},
                {
                    "role": "user",
                    "content": f"""
Job Description:
{st.session_state.job_description}

Resume:
{st.session_state.resume}
"""
                }
            ]
        )

        st.session_state.cover_letter = response.choices[0].message.content


def clear():
    st.session_state.job_description = ""
    st.session_state.resume = ""
    st.session_state.cover_letter = ""

defaults = {
    "job_description": "",
    "resume": "",
    "cover_letter": "",
    "system_prompt": config["system_prompt"],
    "model": config["default_model"]
}

for k, v in defaults.items():
    if k not in st.session_state:
        st.session_state[k] = v

with st.sidebar:
    st.header("Settings")

    st.selectbox(
        "Model",
        options=config["models"],
        key="model"
    )

    st.text_area(
        "System Prompt",
        height=200,
        key="system_prompt"
    )

    st.caption("Tip: Keep prompts strict to avoid generic output.")

col1, col2 = st.columns(2)

with col1:
    st.subheader("Input")

    st.text_area("Job Description", height=250, key="job_description")
    st.text_area("Resume", height=250, key="resume")

    c1, c2 = st.columns(2)

    with c1:
        st.button("Generate", use_container_width=True, on_click=generate)

    with c2:
        st.button("Clear", use_container_width=True, on_click=clear)

with col2:
    st.subheader("Output")

    st.text_area("Cover Letter", height=520, key="cover_letter")

    if st.session_state.cover_letter:
        col_d1, col_d2, col_d3 = st.columns(3)

        with col_d1:
            st.download_button(
                "Download TXT",
                data=st.session_state.cover_letter,
                file_name="cover_letter.txt",
                use_container_width=True
            )

        with col_d2:
            docx_buffer = create_docx_bytes(st.session_state.cover_letter)
            st.download_button(
                "Download DOCX",
                data=docx_buffer,
                file_name="cover_letter.docx",
                use_container_width=True
            )

        with col_d3:
            pdf_buffer = create_pdf_bytes(st.session_state.cover_letter)
            st.download_button(
                "Download PDF",
                data=pdf_buffer,
                file_name="cover_letter.pdf",
                use_container_width=True
            )